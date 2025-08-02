import os
from flask import Flask, render_template, request, send_file
import pdfplumber
import docx
from werkzeug.utils import secure_filename
from fpdf import FPDF
import google.generativeai as genai

# Import analyzer blueprint
try:
    import sys
    import os
    sys.path.append(os.path.dirname(os.path.abspath(__file__)))
    from backend.routes.analyzer import analyzer_bp
    ANALYZER_AVAILABLE = True
    ANALYZER_TYPE = "full"
except ImportError:
    try:
        # Fallback to mock analyzer
        from analyzer_api_mock import analyzer_bp
        ANALYZER_AVAILABLE = True
        ANALYZER_TYPE = "mock"
        print("Using mock Answer Analyzer API for demonstration")
    except ImportError as e:
        ANALYZER_AVAILABLE = False
        ANALYZER_TYPE = "none"
        print(f"Warning: Answer Analyzer module not available: {e}")

# Flask app setup
app = Flask(__name__)
app.config['UPLOAD_FOLDER'] = 'uploads/'
app.config['RESULTS_FOLDER'] = 'results/'
app.config['ALLOWED_EXTENSIONS'] = {'pdf', 'txt', 'docx'}
app.config['MAX_CONTENT_LENGTH'] = 16 * 1024 * 1024  # 16MB max file size

# Register blueprints
if ANALYZER_AVAILABLE:
    app.register_blueprint(analyzer_bp)
    print(f"✅ Answer Analyzer API endpoints registered ({ANALYZER_TYPE} version)")

# Ensure directories exist
os.makedirs(app.config['UPLOAD_FOLDER'], exist_ok=True)
os.makedirs(app.config['RESULTS_FOLDER'], exist_ok=True)

# Initialize Google Gemini API
genai.configure(api_key="AIzaSyBLgzVEI2s9iBPIvswRkA-i6WRTnLMP1FE")
model = genai.GenerativeModel('gemini-2.5-pro')

# Helper function to generate content with Gemini
def generate_with_gemini(prompt):
    try:
        response = model.generate_content(prompt)
        return response.text
    except Exception as e:
        return f"Error generating content: {str(e)}"

# Helper functions
def allowed_file(filename):
    return '.' in filename and filename.rsplit('.', 1)[1].lower() in app.config['ALLOWED_EXTENSIONS']

def extract_text_from_file(file_path):
    text = ""
    if file_path.endswith('.pdf'):
        try:
            with pdfplumber.open(file_path) as pdf:
                for page in pdf.pages:
                    text += page.extract_text() or ""
        except Exception as e:
            return f"Error reading PDF: {str(e)}"
    elif file_path.endswith('.docx'):
        try:
            doc = docx.Document(file_path)
            for paragraph in doc.paragraphs:
                text += paragraph.text + "\n"
        except Exception as e:
            return f"Error reading DOCX: {str(e)}"
    elif file_path.endswith('.txt'):
        try:
            with open(file_path, 'r', encoding='utf-8') as file:
                text = file.read()
        except Exception as e:
            return f"Error reading TXT: {str(e)}"
    return text

def save_mcqs_to_file(content, filename):
    file_path = os.path.join(app.config['RESULTS_FOLDER'], filename)
    try:
        with open(file_path, 'w', encoding='utf-8') as file:
            file.write(content)
    except Exception as e:
        return f"Error saving file: {str(e)}"

def create_pdf(content, filename):
    file_path = os.path.join(app.config['RESULTS_FOLDER'], filename)
    try:
        pdf = FPDF()
        pdf.add_page()
        pdf.set_font("Arial", size=12)
        
        # Split content into lines and add to PDF
        lines = content.split('\n')
        for line in lines:
            # Handle long lines by wrapping them
            if len(line) > 80:
                words = line.split(' ')
                current_line = ""
                for word in words:
                    if len(current_line + word) < 80:
                        current_line += word + " "
                    else:
                        pdf.cell(200, 10, txt=current_line.encode('latin-1', 'replace').decode('latin-1'), ln=True, align='L')
                        current_line = word + " "
                if current_line:
                    pdf.cell(200, 10, txt=current_line.encode('latin-1', 'replace').decode('latin-1'), ln=True, align='L')
            else:
                pdf.cell(200, 10, txt=line.encode('latin-1', 'replace').decode('latin-1'), ln=True, align='L')
        
        pdf.output(file_path)
    except Exception as e:
        return f"Error creating PDF: {str(e)}"

# MCQ generation
def generate_mcqs_with_gemini(text, num_questions):
    prompt = f"""
You are an AI assistant helping the user generate multiple-choice questions (MCQs) from the text below:

Text:
{text}

Generate {num_questions} MCQs. Each should include:
- A clear question
- Four answer options labeled A, B, C, and D
- The correct answer clearly indicated at the end

Format:
## MCQ
Question: [question]
A) [option A]
B) [option B]
C) [option C]
D) [option D]
Correct Answer: [correct option]
"""
    return generate_with_gemini(prompt)

def generate_summary_with_gemini(text):
    prompt = f"""
You are an AI assistant. Read the following text and create a concise summary or study notes that capture the key points and important information.

Text:
{text}

Summary/Study Notes:
"""
    return generate_with_gemini(prompt)

def generate_flashcards_with_gemini(text, num_flashcards):
    prompt = f"""
You are an AI assistant. Read the following text and generate {num_flashcards} flash cards. 
Each flash card should have a clear question (front) and a concise answer (back).

Text:
{text}

Format:
## Flash Card
Q: [question]
A: [answer]
"""
    return generate_with_gemini(prompt)

def generate_long_answers_with_gemini(text, num_questions):
    prompt = f"""
You are an AI assistant. Read the following text and generate {num_questions} long-answer questions. 
For each, provide a detailed answer suitable for exams.

Text:
{text}

Format:
## Long Answer
Question: [question]
Answer: [detailed answer]
"""
    return generate_with_gemini(prompt)

@app.route('/generate', methods=['POST'])
def generate_mcqs():
    if 'file' not in request.files:
        return "No file uploaded."

    file = request.files['file']
    if file and allowed_file(file.filename):
        filename = secure_filename(file.filename)
        file_path = os.path.join(app.config['UPLOAD_FOLDER'], filename)
        file.save(file_path)

        text = extract_text_from_file(file_path)
        if text:
            task = request.form.get('task', 'mcq')
            base_name = filename.rsplit('.', 1)[0]

            if task == 'mcq':
                num_questions = int(request.form['num_questions'])
                mcqs = generate_mcqs_with_gemini(text, num_questions)
                txt_file = f"generated_mcqs_{base_name}.txt"
                pdf_file = f"generated_mcqs_{base_name}.pdf"
                save_mcqs_to_file(mcqs, txt_file)
                create_pdf(mcqs, pdf_file)
                return render_template('results.html', mcqs=mcqs, txt_filename=txt_file, pdf_filename=pdf_file)
            elif task == 'summary':
                summary = generate_summary_with_gemini(text)
                txt_file = f"summary_{base_name}.txt"
                pdf_file = f"summary_{base_name}.pdf"
                save_mcqs_to_file(summary, txt_file)
                create_pdf(summary, pdf_file)
                return render_template('results.html', mcqs=summary, txt_filename=txt_file, pdf_filename=pdf_file)
            elif task == 'flashcards':
                num_flashcards = int(request.form.get('num_questions', 5))
                flashcards = generate_flashcards_with_gemini(text, num_flashcards)
                txt_file = f"flashcards_{base_name}.txt"
                pdf_file = f"flashcards_{base_name}.pdf"
                save_mcqs_to_file(flashcards, txt_file)
                create_pdf(flashcards, pdf_file)
                return render_template('results.html', mcqs=flashcards, txt_filename=txt_file, pdf_filename=pdf_file)
            elif task == 'longanswers':
                num_questions = int(request.form.get('num_questions', 5))
                long_answers = generate_long_answers_with_gemini(text, num_questions)
                txt_file = f"long_answers_{base_name}.txt"
                pdf_file = f"long_answers_{base_name}.pdf"
                save_mcqs_to_file(long_answers, txt_file)
                create_pdf(long_answers, pdf_file)
                return render_template('results.html', mcqs=long_answers, txt_filename=txt_file, pdf_filename=pdf_file)

    return "Invalid file format or upload error."

@app.route('/')
def index():
    return render_template('index.html')

@app.route('/download/<filename>')
def download_file(filename):
    file_path = os.path.join(app.config['RESULTS_FOLDER'], filename)
    return send_file(file_path, as_attachment=True)

if __name__ == '__main__':
    app.run(debug=True)
